import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Load the existing output file
output_file = 'Washout_Matrix_Output.xlsx'

print("Loading existing output file...")
df = pd.read_excel(output_file, sheet_name='Washout Matrix')

# Create simplified sheet with just From Product, To Product, and Duration
simple_df = df[['From_Product', 'To_Product', 'Duration_Hours']].copy()
simple_df.columns = ['From Product', 'To Product', 'Duration (Hours)']

print(f"Creating simplified sheet with {len(simple_df)} transitions...")

# Load workbook and add the new sheet
wb = openpyxl.load_workbook(output_file)

# Remove the sheet if it already exists
if 'Simple Matrix' in wb.sheetnames:
    del wb['Simple Matrix']

# Create new sheet
ws = wb.create_sheet('Simple Matrix', 0)  # Insert as first sheet

# Write headers
ws['A1'] = 'From Product'
ws['B1'] = 'To Product'
ws['C1'] = 'Duration (Hours)'

# Write data
for idx, row in simple_df.iterrows():
    ws.cell(row=idx+2, column=1, value=row['From Product'])
    ws.cell(row=idx+2, column=2, value=row['To Product'])
    ws.cell(row=idx+2, column=3, value=row['Duration (Hours)'])

# Format the sheet
# Header formatting
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(bold=True, color="FFFFFF", size=11)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

# Auto-adjust column widths
for column in ws.columns:
    max_length = 0
    column_letter = get_column_letter(column[0].column)
    
    for cell in column:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    
    adjusted_width = min(max_length + 2, 50)
    ws.column_dimensions[column_letter].width = adjusted_width

# Freeze header row
ws.freeze_panes = 'A2'

# Save the workbook
wb.save(output_file)
print(f"\nSuccessfully added 'Simple Matrix' sheet to {output_file}")
print(f"Sheet contains {len(simple_df)} rows with From Product, To Product, and Duration")

# Show sample
print("\nSample data (first 10 rows):")
print(simple_df.head(10).to_string(index=False))

print("\n" + "="*60)
print("COMPLETE!")
print("="*60)
