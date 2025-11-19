"""
Compare Python-generated output with original Excel washout matrix
to verify the logic is correctly replicated
"""

import pandas as pd
import openpyxl

print("=" * 80)
print("Comparing Python Output vs Original Excel Data")
print("=" * 80)

# Load Python-generated output
python_output = pd.read_excel('Washout_Matrix_Output.xlsx', 
                               sheet_name='Washout Matrix')

# Load original Excel data
wb = openpyxl.load_workbook('Alkox Washout Matrix Product Wheel.xlsm', 
                             data_only=True)
ws = wb['Sheet1']

# Read original data (columns C-I, rows 3 onwards)
original_data = []
for row in range(3, 9219):  # 9216 rows of data
    from_prod = ws.cell(row=row, column=3).value
    if not from_prod:
        break
    
    original_data.append({
        'From_Product': from_prod,
        'To_Product': ws.cell(row=row, column=4).value,
        'Washout': ws.cell(row=row, column=5).value,
        'Washout_Type': ws.cell(row=row, column=6).value,
        'Acid_Wash': ws.cell(row=row, column=7).value,
        'Duration': ws.cell(row=row, column=9).value
    })

wb.close()

original_df = pd.DataFrame(original_data)

print(f"\nOriginal Excel rows: {len(original_df)}")
print(f"Python generated rows: {len(python_output)}")

# Compare matching rows
print("\n" + "=" * 80)
print("Sample Comparison (first 20 transitions)")
print("=" * 80)

for i in range(min(20, len(original_df))):
    orig = original_df.iloc[i]
    
    # Find matching row in Python output
    match = python_output[
        (python_output['From_Product'] == orig['From_Product']) &
        (python_output['To_Product'] == orig['To_Product'])
    ]
    
    if len(match) > 0:
        py = match.iloc[0]
        
        # Map Excel values to Python values
        orig_washout = "YES" if orig['Washout'] == "YES" else "NO"
        py_washout = py['Washout_Needed']
        
        orig_acid = "YES" if orig['Acid_Wash'] == "YES" else "NO"
        py_acid = py['Acid_Wash']
        
        washout_match = "OK" if orig_washout == py_washout else "DIFF"
        type_match = "OK" if orig['Washout_Type'] == py['Washout_Type'] or (orig['Washout_Type'] is None and py['Washout_Type'] != py['Washout_Type']) else "DIFF"
        acid_match = "OK" if orig_acid == py_acid else "DIFF"
        
        print(f"\nRow {i+1}: {orig['From_Product'][:25]:25} -> {orig['To_Product'][:25]:25}")
        print(f"  Washout: {orig_washout:3} vs {py_washout:3} {washout_match}")
        print(f"  Type:    {str(orig['Washout_Type'])[:15]:15} vs {str(py['Washout_Type'])[:15]:15} {type_match}")
        print(f"  Acid:    {orig_acid:3} vs {py_acid:3} {acid_match}")

# Overall statistics comparison
print("\n" + "=" * 80)
print("Overall Statistics Comparison")
print("=" * 80)

print("\nOriginal Excel:")
print(f"  YES washout: {(original_df['Washout'] == 'YES').sum()}")
print(f"  NO washout: {(original_df['Washout'] == 'NO').sum()}")
print(f"  HEEL WASH: {(original_df['Washout_Type'] == 'HEEL WASH').sum()}")
print(f"  FULL WASH: {(original_df['Washout_Type'] == 'FULL WASH').sum()}")
print(f"  Acid wash YES: {(original_df['Acid_Wash'] == 'YES').sum()}")

print("\nPython Generated:")
print(f"  YES washout: {(python_output['Washout_Needed'] == 'YES').sum()}")
print(f"  NO washout: {(python_output['Washout_Needed'] == 'NO').sum()}")
print(f"  HEEL WASH: {(python_output['Washout_Type'] == 'HEEL WASH').sum()}")
print(f"  FULL WASH: {(python_output['Washout_Type'] == 'FULL WASH').sum()}")
print(f"  Acid wash YES: {(python_output['Acid_Wash'] == 'YES').sum()}")

# Check for exact matches on a subset
print("\n" + "=" * 80)
print("Detailed Match Analysis (first 100 rows)")
print("=" * 80)

matches = 0
washout_matches = 0
type_matches = 0
acid_matches = 0

for i in range(min(100, len(original_df))):
    orig = original_df.iloc[i]
    match = python_output[
        (python_output['From_Product'] == orig['From_Product']) &
        (python_output['To_Product'] == orig['To_Product'])
    ]
    
    if len(match) > 0:
        py = match.iloc[0]
        matches += 1
        
        orig_washout = "YES" if orig['Washout'] == "YES" else "NO"
        if orig_washout == py['Washout_Needed']:
            washout_matches += 1
        
        # For type, both should be None/NaN or equal
        orig_type = orig['Washout_Type']
        py_type = py['Washout_Type']
        if (pd.isna(orig_type) and pd.isna(py_type)) or (orig_type == py_type):
            type_matches += 1
        
        orig_acid = "YES" if orig['Acid_Wash'] == "YES" else "NO"
        if orig_acid == py['Acid_Wash']:
            acid_matches += 1

print(f"\nOut of 100 rows:")
print(f"  Product pairs matched: {matches}/100")
print(f"  Washout decision matches: {washout_matches}/100 ({washout_matches}%)")
print(f"  Washout type matches: {type_matches}/100 ({type_matches}%)")
print(f"  Acid wash matches: {acid_matches}/100 ({acid_matches}%)")

if washout_matches == 100 and type_matches == 100 and acid_matches == 100:
    print("\n[OK] Perfect match! Python logic matches Excel exactly.")
else:
    print(f"\n[NOTICE] Some differences found. This may be expected if the original Excel")
    print(f"  uses different formulas or has been manually edited.")

print("\n" + "=" * 80)
